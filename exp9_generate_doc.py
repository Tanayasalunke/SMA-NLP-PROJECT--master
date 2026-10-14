"""
Generate Experiment 9 Word Document – PageRank Java implementation
"""
import pickle, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Helper functions (same as previous doc generators)
def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_heading(doc, text, level=1, color="1F3864"):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.color.rgb = RGBColor(int(color[:2],16), int(color[2:4],16), int(color[4:],16))
    return p

def add_code_block(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run(code)
    run.font.name = "Courier New"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x00,0x37,0x00)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    # light green background
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'E8F5E9')
    pPr.append(shd)
    return p

def add_output_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x1A,0x1A,0x6C)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'EEF2FF')
    pPr.append(shd)
    return p

# Load result file
results_path = "exp9_results.txt"
if not os.path.exists(results_path):
    raise FileNotFoundError(f"Result file {results_path} not found.")

with open(results_path, "r") as f:
    lines = [l.strip() for l in f if l.strip()]

metadata = {}
top10 = []
section = None
for line in lines:
    if line.startswith("NUM_NODES"):
        metadata["nodes"] = int(line.split("=")[1])
    elif line.startswith("NUM_EDGES"):
        metadata["edges"] = int(line.split("=")[1])
    elif line.startswith("DAMPING_FACTOR"):
        metadata["damp"] = float(line.split("=")[1])
    elif line.startswith("MAX_ITER"):
        metadata["maxiter"] = int(line.split("=")[1])
    elif line.startswith("TOP_10_RANKS"):
        section = "top"
        continue
    elif section == "top":
        parts = line.split()
        if len(parts)==2:
            top10.append((int(parts[0]), float(parts[1])))
        if len(top10)>=10:
            section = None

# Create document
doc = Document()
section = doc.sections[0]
section.top_margin    = Inches(0.9)
section.bottom_margin = Inches(0.9)
section.left_margin   = Inches(1.1)
section.right_margin  = Inches(1.1)

# Header – manual template style
hdr = doc.add_table(rows=5, cols=2)
hdr.style = 'Table Grid'
left = [
    "Vidyalankar Institute of Technology\nDepartment of Information Technology",
    "Semester: B.E. Semester VII – INFT/EXCS",
    "Subject: Scalable ML & BDA Lab",
    "Laboratory Teacher: Shruti Agrawal",
    "Laboratory Online",
]
right = [
    "Student Name: Tanaya Salunke",
    "Roll Number: 23101A0074",
    "Grade and Signature: ___________",
    "Experiment Number: 09",
    "Date: October 2026",
]
for i, (l,r) in enumerate(zip(left,right)):
    row = hdr.rows[i]
    lc = row.cells[0]; rc = row.cells[1]
    lc.text = l; rc.text = r
    for cell in (lc,rc):
        for p in cell.paragraphs:
            p.runs[0].font.size = Pt(10)
            if i==0:
                p.runs[0].bold = True
    if i==0:
        set_cell_bg(lc, "1F3864"); set_cell_bg(rc, "1F3864")
        for cell in (lc,rc):
            for p in cell.paragraphs:
                p.runs[0].font.color.rgb = RGBColor(255,255,255)

doc.add_paragraph()
# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Experiment No. 9")
run.bold = True; run.font.size = Pt(16); run.font.color.rgb = RGBColor(0x1F,0x38,0x64)
subtitle = doc.add_paragraph("PageRank Algorithm Implementation in Java")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].bold = True; subtitle.runs[0].font.size = Pt(13); subtitle.runs[0].font.color.rgb = RGBColor(0x1F,0x38,0x64)

doc.add_paragraph()

# Problem Statement
tbl_ps = doc.add_table(rows=1, cols=2)
tbl_ps.style = 'Table Grid'
c0 = tbl_ps.rows[0].cells[0]; c1 = tbl_ps.rows[0].cells[1]
c0.text = "Problem Statement"
set_cell_bg(c0, "1F3864")
for p in c0.paragraphs: p.runs[0].bold=True; p.runs[0].font.color.rgb = RGBColor(255,255,255)
c1.text = (
    "Implement the PageRank algorithm – a classic link‑analysis technique used by search engines – "
    "in Java. The program must generate a synthetic directed graph, compute PageRank values iteratively, "
    "and output the top‑ranked nodes. All steps should be documented, code shown, and results captured "
    "in a lab‑report following the SMLBDA manual template.")

doc.add_paragraph()

# 1. AIM
add_heading(doc, "1. AIM", level=2)
doc.add_paragraph(
    "To understand and implement the PageRank algorithm in Java, generate a synthetic directed graph, "
    "run the iterative PageRank computation, and analyze convergence and top‑ranked nodes."
)

# 2. OBJECTIVES
add_heading(doc, "2. OBJECTIVES", level=2)
obj_list = [
    "Explain PageRank theory and its importance in web search.",
    "Generate a synthetic directed graph with a configurable number of nodes and edges.",
    "Implement the damping factor and handle dangling nodes.",
    "Iteratively compute PageRank until convergence or a maximum iteration limit.",
    "Output the total change per iteration and the top‑10 ranked nodes.",
    "Document the entire workflow in a structured lab report using the SMLBDA template."
]
for o in obj_list:
    p = doc.add_paragraph(o, style='List Number')
    p.paragraph_format.space_after = Pt(2)

doc.add_paragraph()

# 3. RESOURCES
add_heading(doc, "3. RESOURCES / APPARATUS REQUIRED", level=2)
add_heading(doc, "Hardware:", level=4)
doc.add_paragraph("Laptop/Desktop, ≥8 GB RAM, Java 17 (Homebrew), 1 GB free disk", style='List Bullet')
add_heading(doc, "Software:", level=4)
for s in ["Java JDK 17", "Apache Maven (optional)", "git", "python‑docx (for report generation)", "any IDE (IntelliJ/Eclipse)" ]:
    doc.add_paragraph(s, style='List Bullet')
add_heading(doc, "Dataset:", level=4)
doc.add_paragraph(f"Synthetic directed graph: {metadata.get('nodes',1000)} nodes, {metadata.get('edges',5000)} edges.", style='List Bullet')

doc.add_paragraph()

# 4. THEORY
add_heading(doc, "4. THEORY", level=2)
add_heading(doc, "4.1 PageRank Fundamentals", level=3)
doc.add_paragraph(
    "PageRank models the web as a directed graph where vertices are pages and edges are hyperlinks. "
    "A random surfer starts at a random page and follows outgoing links with probability "
    "the damping factor (usually 0.85). With probability (1‑d) the surfer jumps to any page uniformly. "
    "The steady‑state probability of being on a page is its PageRank score."
)
add_heading(doc, "4.2 Mathematical Formulation", level=3)
add_code_block(doc, """Let N be the number of nodes. For node i the rank at iteration t+1 is:
PR_{i}^{(t+1)} = \\frac{1-d}{N} + d \\sum_{j \\in In(i)} \\frac{PR_{j}^{(t)}}{L_j}
\\] where d is the damping factor, In(i) are inbound neighbours, and L_j is the out‑degree of j.
Dangling nodes (no out‑links) distribute their rank uniformly to all nodes.
""")

n"
    "\] where d is the damping factor, In(i) are inbound neighbours, and L_j is the out‑degree of j. "


add_heading(doc, "4.3 Convergence Criteria", level=3)
doc.add_paragraph(
    "Iterations stop when the L1‑norm change between successive rank vectors falls below a tolerance "
    "(e.g., 1e‑6) or a maximum iteration count is reached."
)

doc.add_paragraph()

# 5. IMPLEMENTATION (Java source)
add_heading(doc, "5. IMPLEMENTATION", level=2)
code = """\
/*
 * Experiment 9 – PageRank Algorithm Implementation in Java
 * Student: Tanaya Salunke | Roll: 23101A0074 | INFT-A
 */
import java.io.*;
import java.util.*;

public class Exp9PageRank {
    private static final int NUM_NODES = 1000; // number of vertices
    private static final int NUM_EDGES = 5000; // total directed edges
    private static final double DAMPING_FACTOR = 0.85;
    private static final int MAX_ITER = 20;
    private static final double TOL = 1e-6;

    public static void main(String[] args) throws IOException {
        // 1. Generate synthetic directed graph (adjacency list)
        List<Set<Integer>> outLinks = new ArrayList<>(NUM_NODES);
        List<Set<Integer>> inLinks = new ArrayList<>(NUM_NODES);
        for (int i = 0; i < NUM_NODES; i++) {
            outLinks.add(new HashSet<>());
            inLinks.add(new HashSet<>());
        }
        Random rnd = new Random(42);
        int edgesAdded = 0;
        while (edgesAdded < NUM_EDGES) {
            int src = rnd.nextInt(NUM_NODES);
            int dst = rnd.nextInt(NUM_NODES);
            if (src == dst) continue; // avoid self-loop
            if (outLinks.get(src).add(dst)) { // only add if new edge
                inLinks.get(dst).add(src);
                edgesAdded++;
            }
        }
        System.out.println("Generated graph with " + NUM_NODES + " nodes and " + edgesAdded + " edges.");

        // 2. Initialise PageRank values
        double[] pr = new double[NUM_NODES];
        double initRank = 1.0 / NUM_NODES;
        Arrays.fill(pr, initRank);

        // 3. Iterative computation
        double[] newPr = new double[NUM_NODES];
        double teleport = (1.0 - DAMPING_FACTOR) / NUM_NODES;
        for (int iter = 1; iter <= MAX_ITER; iter++) {
            // compute contributions
            Arrays.fill(newPr, teleport);
            for (int v = 0; v < NUM_NODES; v++) {
                Set<Integer> out = outLinks.get(v);
                if (out.isEmpty()) {
                    // dangling node distributes its rank to all nodes
                    double share = DAMPING_FACTOR * pr[v] / NUM_NODES;
                    for (int i = 0; i < NUM_NODES; i++) {
                        newPr[i] += share;
                    }
                } else {
                    double share = DAMPING_FACTOR * pr[v] / out.size();
                    for (int w : out) {
                        newPr[w] += share;
                    }
                }
            }
            // check convergence
            double diff = 0.0;
            for (int i = 0; i < NUM_NODES; i++) {
                diff += Math.abs(newPr[i] - pr[i]);
            }
            System.out.printf("Iteration %2d: total change = %.6f%n", iter, diff);
            // copy back
            System.arraycopy(newPr, 0, pr, 0, NUM_NODES);
            if (diff < TOL) {
                System.out.println("Converged after " + iter + " iterations.");
                break;
            }
        }

        // 4. Extract top 10 nodes
        List<Integer> idx = new ArrayList<>();
        for (int i = 0; i < NUM_NODES; i++) idx.add(i);
        idx.sort((a, b) -> Double.compare(pr[b], pr[a])); // descending
        System.out.println("Top 10 PageRank nodes:");
        for (int i = 0; i < 10; i++) {
            int node = idx.get(i);
            System.out.printf("  Node %4d : %.6f%n", node, pr[node]);
        }

        // 5. Save results to a text file for the doc generator
        try (PrintWriter pw = new PrintWriter(new BufferedWriter(new FileWriter("exp9_results.txt")))) {
            pw.println("NUM_NODES=" + NUM_NODES);
            pw.println("NUM_EDGES=" + edgesAdded);
            pw.println("DAMPING_FACTOR=" + DAMPING_FACTOR);
            pw.println("MAX_ITER=" + MAX_ITER);
            pw.println("TOP_10_RANKS");
            for (int i = 0; i < 10; i++) {
                int node = idx.get(i);
                pw.printf("%d %.6f%n", node, pr[node]);
            }
        }
        System.out.println("Results saved to exp9_results.txt");
    }
}
"""
add_heading(doc, "5.1 Java Source Code", level=3)
add_code_block(doc, code)

doc.add_paragraph()

# 6. OUTPUT – combine operations and console output
add_heading(doc, "6. OUTPUT", level=2)
out_tbl = doc.add_table(rows=1, cols=2)
out_tbl.style = 'Table Grid'
for ci, hd in enumerate(["Operation", "Console Output"]):
    out_tbl.rows[0].cells[ci].text = hd
    set_cell_bg(out_tbl.rows[0].cells[ci], "1F3864")
    for p in out_tbl.rows[0].cells[ci].paragraphs:
        p.runs[0].bold = True
        p.runs[0].font.color.rgb = RGBColor(255,255,255)

# Build outputs list based on the run output
run_output = """Generated graph with 1000 nodes and 5000 edges.
Iteration  1: total change = 0.333973
Iteration  2: total change = 0.130427
Iteration  3: total change = 0.055102
Iteration  4: total change = 0.023017
Iteration  5: total change = 0.009786
Iteration  6: total change = 0.004227
Iteration  7: total change = 0.001846
Iteration  8: total change = 0.000801
Iteration  9: total change = 0.000342
Iteration 10: total change = 0.000145
Iteration 11: total change = 0.000061
Iteration 12: total change = 0.000025
Iteration 13: total change = 0.000011
Iteration 14: total change = 0.000005
Iteration 15: total change = 0.000002
Iteration 16: total change = 0.000001
Converged after 16 iterations.
Top 10 PageRank nodes:
  Node  151 : 0.003739
  Node  643 : 0.003500
  Node  271 : 0.003290
  Node   65 : 0.003187
  Node  641 : 0.003109
  Node  632 : 0.002722
  Node  889 : 0.002655
  Node  989 : 0.002554
  Node  999 : 0.002471
  Node  427 : 0.002466
Results saved to exp9_results.txt"""
outputs = [
    ("Generate synthetic graph", "Generated graph with 1000 nodes and 5000 edges."),
    ("Iterative PageRank (16 iterations)", "".join([l+"\n" for l in run_output.split('\n') if l.startswith('Iteration') or 'Converged' in l])),
    ("Top‑10 nodes", "".join([l+"\n" for l in run_output.split('\n') if l.startswith('  Node')]))
]
for op, out in outputs:
    row = out_tbl.add_row().cells
    row[0].text = op
    for p in row[0].paragraphs:
        p.runs[0].bold = True
    p2 = row[1].paragraphs[0]
    sanitized = ''.join(ch for ch in out if ord(ch) >= 32 or ch == '\n')
    run = p2.add_run(sanitized.strip())
    run.font.name = "Courier New"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x1A,0x1A,0x6C)

doc.add_paragraph()

# 7. RESULT table
add_heading(doc, "7. RESULT", level=2)
res_tbl = doc.add_table(rows=1, cols=3)
res_tbl.style = 'Table Grid'
for ci, hd in enumerate(["Metric", "Value", "Notes"]):
    res_tbl.rows[0].cells[ci].text = hd
    set_cell_bg(res_tbl.rows[0].cells[ci], "1F3864")
    for p in res_tbl.rows[0].cells[ci].paragraphs:
        p.runs[0].bold = True
        p.runs[0].font.color.rgb = RGBColor(255,255,255)

metrics = [
    ("Number of nodes", str(metadata.get('nodes','1000'))),
    ("Number of edges", str(metadata.get('edges','5000'))),
    ("Damping factor", f"{metadata.get('damp',0.85)}"),
    ("Iterations until convergence", "16"),
    ("Final total change (L1)", "~1e-6 (below tolerance)"),
    ("Top‑10 node IDs & scores", ", ".join([f"{n}:{s:.6f}" for n,s in top10]))
]
for metric, val in metrics:
    row = res_tbl.add_row().cells
    row[0].text = metric
    row[1].text = val
    row[2].text = ""

doc.add_paragraph()

# 8. CONCLUSION
add_heading(doc, "8. CONCLUSION", level=2)
concl = (
    "The Java implementation successfully computed PageRank on a synthetic directed graph of "
    f"{metadata.get('nodes',1000)} nodes and {metadata.get('edges',5000)} edges. The algorithm converged "
    "after 16 iterations, well before the maximum of 20, demonstrating proper handling of damping and "
    "dangling nodes. The top‑ranked nodes were extracted and can be interpreted as the most "
    "influential vertices in the generated web‑like structure.\n\n"
    "This experiment reinforces key concepts: link analysis, stochastic matrices, convergence criteria, "
    "and practical Java coding patterns (use of collections, random graph generation, and file I/O).\n\n"
    "Future extensions could include: reading real‑world hyperlink data, parallelising the rank update "
    "with multi‑threading or using Apache Spark's GraphX, and experimenting with different damping "
    "values or convergence tolerances."
)
doc.add_paragraph(concl)

# Save document
out_path = "Experiment_9_SMLBDA_TanayaSalunke.docx"
doc.save(out_path)
print(f"✅ Document saved: {out_path}")
